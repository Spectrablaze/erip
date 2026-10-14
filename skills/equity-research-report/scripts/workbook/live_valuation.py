"""live_valuation.py - turn the financial-model workbook into the client model.

The financial-model skill already writes a live three-statement workbook: history is data,
every forecast cell is a formula off the Drivers sheet. This adds the valuation on top of it,
also as formulas, so nothing a client sees is a pasted number:

  Scenarios    one selector cell; each case's driver paths as inputs; the active path feeds
               the Drivers sheet, so the statements, checks and DCF all switch together
  WACC         CAPM build (Blume or peer-relevered beta), post-tax cost of debt, WACC
  DCF          FCFF = EBIT x (1 - t) + D&A - capex - change in NWC, linked to the statements;
               mid-year discounting, Gordon terminal value, exit-multiple cross-check,
               target price and the mechanical rating
  Sensitivity  WACC x terminal growth, and WACC x exit multiple, each cell a formula

    python live_valuation.py --root reports/<T> --out reports/<T>/<Name>_ERIP_Model.xlsx

Reads data/scenarios_3s.json (from scenario_runs.py), each case's three-statement workbook under
data/three_statement/scenarios/<case>/, data/assumptions.json and the report's target settings.
Then recalculate it (lo_recalc.py) so the file carries values as well as formulas.
"""
from __future__ import annotations

import argparse
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.worksheet.datavalidation import DataValidation

BLUE, BLACK, GREEN = "FF0000FF", "FF000000", "FF008000"
PCT, NUM, INR, MULT = '0.0%;(0.0%)', '#,##0.0;(#,##0.0);"-"', '#,##0;(#,##0)', '0.0"x"'
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
THIN = Side(style="thin", color="FF8EA9DB")


def _rate(v):
    """Percent or fraction in, fraction out (assumptions.json mixes both)."""
    if v is None:
        return None
    if isinstance(v, list):
        return [_rate(x) for x in v]
    v = float(v)
    return v / 100 if abs(v) > 1 else v


class Sheet:
    """Small writer that keeps the input / formula / link colour convention in one place."""

    def __init__(self, ws, proto):
        self.ws, self.p = ws, proto

    def title(self, text, sub):
        self.ws["A1"], self.ws["A2"] = text, sub
        self.ws["A1"].font = Font(**{**_font(self.p["A1"]), })
        self.ws["A2"].font = Font(**_font(self.p["A2"]))

    def header(self, row, labels, first_col=2, label="INR cr unless stated"):
        cells = [(1, label)] + [(first_col + i, t) for i, t in enumerate(labels)]
        last = max(c for c, _ in cells)
        for c in range(1, last + 1):
            cell = self.ws.cell(row, c)
            cell.fill = PatternFill("solid", fgColor=self.p["A4"].fill.fgColor.rgb)
            cell.font = Font(**_font(self.p["A4"]))
        for c, t in cells:
            self.ws.cell(row, c, t).alignment = Alignment(horizontal="left" if c == 1 else "right")

    def section(self, row, text):
        c = self.ws.cell(row, 1, text)
        c.font = Font(**_font(self.p["A5"]))

    def put(self, row, col, value, kind="formula", fmt=NUM, bold=False):
        c = self.ws.cell(row, col, value)
        c.font = Font(name=self.p["A6"].font.name, size=self.p["A6"].font.sz, bold=bold,
                      color={"input": BLUE, "link": GREEN}.get(kind, BLACK))
        c.number_format = fmt
        if kind == "input":
            c.fill = INPUT_FILL
        return c

    def label(self, row, text, indent=0, bold=False, note=None, note_col=None):
        c = self.ws.cell(row, 1, ("    " * indent) + text)
        c.font = Font(name=self.p["A6"].font.name, size=self.p["A6"].font.sz, bold=bold)
        if note and note_col:
            n = self.ws.cell(row, note_col, note)
            n.font = Font(name=self.p["A6"].font.name, size=self.p["A6"].font.sz - 1, italic=True, color="FF595959")


def _font(cell):
    f = cell.font
    return {"name": f.name, "size": f.sz, "bold": f.b, "italic": f.i,
            "color": f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None}


def _rows(ws):
    return {str(ws.cell(r, 1).value).strip(): r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}


def build(root, out):
    S3 = json.load(open(os.path.join(root, "data", "scenarios_3s.json"), encoding="utf-8"))
    A = json.load(open(os.path.join(root, "data", "assumptions.json"), encoding="utf-8"))
    cfg_p = os.path.join(root, "data", "valuation_settings.json")
    cfg = json.load(open(cfg_p, encoding="utf-8")) if os.path.exists(cfg_p) else {}
    cases = list(S3["cases"])
    order = ["base"] + [c for c in ("bull", "bear") if c in cases] + [c for c in cases if c not in ("base", "bull", "bear")]
    ts = lambda c: os.path.join(root, "data", "three_statement", "scenarios", c, "model.xlsx")
    wb = openpyxl.load_workbook(ts("base"))
    other = {c: openpyxl.load_workbook(ts(c)) for c in order if c != "base"}

    D = wb["Drivers"]
    proto = D
    hdr = [D.cell(4, c).value for c in range(1, D.max_column + 1)]
    fc = [c for c in range(2, D.max_column + 1) if str(hdr[c - 1] or "").endswith("E")]
    years = [hdr[c - 1] for c in fc]
    N = len(fc)
    last_hist = fc[0] - 1
    R = {name: _rows(wb[name]) for name in ("Drivers", "Income Statement", "Balance Sheet", "Cash Flow", "Schedules")}

    # ---- which driver rows actually differ between cases ------------------------------------
    varying = []
    for r in range(5, D.max_row + 1):
        if not D.cell(r, 1).value or D.cell(r, fc[0]).value is None:
            continue
        paths = {"base": [D.cell(r, c).value for c in fc]}
        for c, owb in other.items():
            paths[c] = [owb["Drivers"].cell(r, k).value for k in fc]
        if any(paths[c] != paths["base"] for c in paths):
            if any(isinstance(v, str) for p in paths.values() for v in p):
                raise SystemExit(f"Drivers row {r} differs between cases but holds formulas; cannot switch it")
            varying.append((r, str(D.cell(r, 1).value).strip(), paths))

    sc_meta = (A.get("scenarios") or {})
    tg = {c: _rate(sc_meta.get(c, {}).get("terminal_growth", A["terminal_growth"])) for c in order}
    prob = {c: S3["cases"][c]["probability"] for c in order}

    # ---- Scenarios sheet ---------------------------------------------------------------------
    ws = wb.create_sheet("Scenarios", index=wb.sheetnames.index("Drivers"))
    Sc = Sheet(ws, proto)
    Sc.title(f"{D['A1'].value.split(' - ')[0]} - Scenarios",
             "Change the yellow cell: every switched driver, the statements, the checks and the DCF follow")
    ws["A4"] = "Active scenario  (" + ", ".join(f"{i + 1} {c.title()}" for i, c in enumerate(order)) + ")"
    ws["A4"].font = Font(name=proto["A6"].font.name, size=proto["A6"].font.sz, bold=True)
    Sc.put(4, 2, 1, "input", "0")
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(len(order)), allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    Sc.put(5, 2, "=CHOOSE($B$4," + ",".join(f'"{c.title()}"' for c in order) + ")", fmt="@", bold=True)
    Sc.label(5, "Active case")
    SEL = "Scenarios!$B$4"

    row = 7
    Sc.header(row, years)
    row += 1
    blocks = {}
    for c in order:
        why = (sc_meta.get(c) or {}).get("_why", "")
        Sc.section(row, f"{c.title()} case inputs  (probability {prob[c]:.0%})")
        row += 1
        blocks[c] = {}
        for r, lab, paths in varying:
            Sc.label(row, lab, indent=1)
            fmt = D.cell(r, fc[0]).number_format
            for i, v in enumerate(paths[c]):
                Sc.put(row, 2 + i, v, "input", fmt)
            blocks[c][r] = row
            row += 1
        Sc.label(row, "Terminal growth", indent=1)
        Sc.put(row, 2, tg[c], "input", PCT)
        blocks[c]["tg"] = row
        row += 1
        Sc.label(row, "Probability", indent=1)
        Sc.put(row, 2, prob[c], "input", '0%')
        blocks[c]["p"] = row
        row += 1
        if why:
            ws.cell(row, 1, why[:400]).font = Font(size=8, italic=True, color="FF595959")
            ws.cell(row, 1).alignment = Alignment(wrap_text=False)
            row += 1
        row += 1

    Sc.section(row, "Active path  (feeds the Drivers sheet)")
    row += 1
    active = {}
    for r, lab, _ in varying:
        Sc.label(row, lab, indent=1)
        fmt = D.cell(r, fc[0]).number_format
        for i in range(N):
            col = COL(2 + i)
            Sc.put(row, 2 + i, "=CHOOSE(" + SEL + "," + ",".join(f"{col}{blocks[c][r]}" for c in order) + ")", fmt=fmt)
        active[r] = row
        row += 1
    Sc.label(row, "Terminal growth", indent=1)
    Sc.put(row, 2, "=CHOOSE(" + SEL + "," + ",".join(f"B{blocks[c]['tg']}" for c in order) + ")", fmt=PCT)
    TG = f"Scenarios!$B${row}"
    row += 2

    # Recorded per-case values: the only numbers on this sheet that are outputs. Each was
    # produced by recalculating this workbook (and the Python model) with the selector at that case.
    Sc.section(row, "Scenario values  (INR per share; recorded with the selector at each case)")
    row += 1
    Sc.header(row, ["Prob.", "Value", "Weighted"], label="Case")
    row += 1
    first = row
    for c in order:
        Sc.label(row, c.title(), indent=1)
        Sc.put(row, 2, f"=B{blocks[c]['p']}", fmt='0%')
        Sc.put(row, 3, S3["cases"][c]["intrinsic_value_per_share"], "input", INR)
        Sc.put(row, 4, f"=B{row}*C{row}", fmt=INR)
        row += 1
    Sc.label(row, "Probability-weighted value", bold=True)
    Sc.put(row, 2, f"=SUM(B{first}:B{row - 1})", fmt='0%')
    Sc.put(row, 4, f"=SUM(D{first}:D{row - 1})", fmt=INR, bold=True)
    WTD_ROW = row
    row += 2
    Sc.label(row, "Live value, active case (DCF sheet)")
    live_ref_row = row
    row += 1
    Sc.label(row, "Recorded value, active case")
    Sc.put(row, 2, f"=INDEX(C{first}:C{first + len(order) - 1},{SEL})", fmt=INR)
    row += 1
    Sc.label(row, "Difference (must be 0)", bold=True)
    Sc.put(row, 2, f"=ROUND(B{live_ref_row},0)-B{row - 1}", fmt=INR, bold=True)
    ws.column_dimensions["A"].width = 52
    for i in range(N):
        ws.column_dimensions[COL(2 + i)].width = 11
    ws.freeze_panes = "B8"

    # The Drivers rows that switch become links to the active path.
    for r, _, _ in varying:
        for i, c in enumerate(fc):
            cell = D.cell(r, c)
            cell.value = f"=Scenarios!{COL(2 + i)}{active[r]}"
            cell.font = Font(name=cell.font.name, size=cell.font.sz, color=GREEN)
            cell.fill = PatternFill()

    # ---- WACC ----------------------------------------------------------------------------------
    ws = wb.create_sheet("WACC")
    W = Sheet(ws, proto)
    W.title(f"{D['A1'].value.split(' - ')[0]} - Cost of capital", "CAPM cost of equity, post-tax cost of debt, target weights")
    W.header(4, ["Value", "Basis"], label="")
    tax = _rate(A["tax_rate"])
    lines = [("Risk-free rate", _rate(A["rf"]), "input", PCT, "10-year G-sec yield (assumptions.json)"),
             ("Equity risk premium", _rate(A["erp"]), "input", PCT, "assumptions.json"),
             ("Raw beta", float(A["beta"]), "input", "0.00", "assumptions.json")]
    r = 6
    W.section(5, "Cost of equity")
    for lab, v, k, f, note in lines:
        W.label(r, lab, 1, note=note, note_col=3)
        W.put(r, 2, v, k, f)
        r += 1
    rf_r, erp_r, braw_r = 6, 7, 8
    beta_ref = f"B{braw_r}"
    if A.get("peer_betas"):
        W.section(r, "Peer betas (Hamada unlever, median, relever at target D/E)")
        r += 1
        p0 = r
        for p in A["peer_betas"]:
            W.label(r, f"{p.get('name', 'peer')}  beta / D/E / unlevered", 1)
            W.put(r, 2, float(p["beta"]), "input", "0.00")
            W.put(r, 3, float(p["debt_equity"]), "input", "0.00")
            W.put(r, 4, f"=B{r}/(1+(1-B{{TAX}})*C{r})", fmt="0.00")
            r += 1
        unlev = f"MEDIAN(D{p0}:D{r - 1})"
        W.label(r, "Relevered beta at target D/E", 1)
        W.put(r, 2, f"={unlev}*(1+(1-B{{TAX}})*B{{WD}}/(1-B{{WD}}))", fmt="0.00")
        beta_ref = f"B{r}"
        r += 1
    if A.get("blume_adjust"):
        W.label(r, "Blume-adjusted beta (0.67 x beta + 0.33)", 1)
        W.put(r, 2, f"=0.67*{beta_ref}+0.33", fmt="0.00")
        beta_ref = f"B{r}"
        r += 1
    W.label(r, "Cost of equity (CAPM)", bold=True)
    W.put(r, 2, f"=B{rf_r}+{beta_ref}*B{erp_r}", fmt=PCT, bold=True)
    ke_r = r
    r += 2
    W.section(r, "Cost of debt and weights")
    r += 1
    W.label(r, "Pre-tax cost of debt", 1, note="assumptions.json", note_col=3)
    W.put(r, 2, _rate(A.get("cost_of_debt", float(A["rf"]) + 1.5)), "input", PCT)
    kd_r = r
    r += 1
    W.label(r, "Tax rate", 1, note="assumptions.json", note_col=3)
    W.put(r, 2, tax, "input", PCT)
    tax_r = r
    r += 1
    W.label(r, "Post-tax cost of debt", 1)
    W.put(r, 2, f"=B{kd_r}*(1-B{tax_r})", fmt=PCT)
    kdt_r = r
    r += 1
    W.label(r, "Target debt weight", 1, note="assumptions.json", note_col=3)
    W.put(r, 2, _rate(A.get("target_debt_weight", 0)), "input", PCT)
    wd_r = r
    r += 2
    W.label(r, "WACC", bold=True)
    W.put(r, 2, f"=B{ke_r}*(1-B{wd_r})+B{kdt_r}*B{wd_r}", fmt='0.00%', bold=True)
    WACC = f"WACC!$B${r}"
    r += 1
    W.label(r, "Terminal growth (active scenario)", 1)
    W.put(r, 2, f"={TG}", "link", PCT)
    r += 1
    W.label(r, "Terminal growth below WACC", 1)
    W.put(r, 2, f'=IF(B{r - 1}<{WACC},"OK","FAIL")', fmt="@").alignment = Alignment(horizontal="right")
    for row_ in ws.iter_rows(min_row=1, max_row=r):
        for c in row_:
            if isinstance(c.value, str) and "{TAX}" in c.value:
                c.value = c.value.replace("{TAX}", str(tax_r)).replace("{WD}", str(wd_r))
            if isinstance(c.value, str) and "{WD}" in c.value:
                c.value = c.value.replace("{WD}", str(wd_r))
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 44

    # ---- DCF -----------------------------------------------------------------------------------
    ws = wb.create_sheet("DCF")
    F = Sheet(ws, proto)
    F.title(f"{D['A1'].value.split(' - ')[0]} - DCF",
            "FCFF from the linked statements; mid-year discounting; Gordon terminal value; exit-multiple target")
    F.header(4, years)
    IS, SCH, DRV = R["Income Statement"], R["Schedules"], R["Drivers"]

    def link(sheet, rr, i):
        return f"='{sheet}'!{COL(fc[i])}{rr}"

    F.section(5, "Free cash flow to the firm")
    spec = [("Revenue", lambda i: link("Income Statement", IS["Revenue"], i), "link", NUM, False),
            ("EBITDA", lambda i: link("Income Statement", IS["EBITDA"], i), "link", NUM, False),
            ("EBIT", lambda i: link("Income Statement", IS["EBIT"], i), "link", NUM, False),
            ("Tax rate", lambda i: link("Drivers", DRV["Effective tax rate"], i), "link", PCT, False),
            ("NOPAT  (EBIT x (1 - tax rate))", lambda i: f"={COL(2 + i)}8*(1-{COL(2 + i)}9)", "formula", NUM, False),
            ("Add: depreciation and amortisation", lambda i: link("Income Statement", IS["Depreciation and amortisation"], i), "link", NUM, False),
            ("Less: capital expenditure", lambda i: f"=-'Schedules'!{COL(fc[i])}{SCH['Capital expenditure']}", "link", NUM, False),
            ("Less: increase in net operating assets", lambda i: f"=-'Schedules'!{COL(fc[i])}{SCH['Change in net operating assets']}", "link", NUM, False),
            ("Free cash flow to the firm", lambda i: f"={COL(2 + i)}10+{COL(2 + i)}11+{COL(2 + i)}12+{COL(2 + i)}13", "formula", NUM, True)]
    for k, (lab, fn, kind, fmt, bold) in enumerate(spec):
        rr = 6 + k
        F.label(rr, lab, indent=0 if bold or k < 3 else 1, bold=bold)
        for i in range(N):
            F.put(rr, 2 + i, fn(i), kind, fmt, bold)
    FCFF_R = 14
    F.section(16, "Discounting")
    F.label(17, "Discount period (mid-year)", 1)
    F.put(17, 2, 0.5, "input", "0.0")  # mid-year: first year's cash arrives at t = 0.5
    for i in range(1, N):
        F.put(17, 2 + i, f"={COL(1 + i)}17+1", fmt="0.0")
    F.label(18, "Discount factor", 1)
    F.label(19, "Present value of FCFF", 1)
    for i in range(N):
        F.put(18, 2 + i, f"=1/(1+{WACC})^{COL(2 + i)}17", fmt="0.000")
        F.put(19, 2 + i, f"={COL(2 + i)}14*{COL(2 + i)}18", fmt=NUM)
    LASTC = COL(1 + N)
    FIRSTC = "B"

    nd = float(A.get("net_debt", 0))
    sh_ref = f"='Schedules'!{COL(last_hist)}{SCH['Shares outstanding']}"
    rows = [
        ("Valuation: perpetuity growth", None, None, None, None),
        ("WACC", f"={WACC}", "link", '0.00%', None),
        ("Terminal growth (active scenario)", f"={TG}", "link", PCT, None),
        ("Sum of PV of FCFF", f"=SUM({FIRSTC}19:{LASTC}19)", "formula", NUM, None),
        ("Terminal value  (FCFF N x (1 + g) / (WACC - g))", f"={LASTC}14*(1+B23)/(B22-B23)", "formula", NUM, None),
        ("PV of terminal value  (discounted at N - 0.5)", f"=B25/(1+B22)^{LASTC}17", "formula", NUM, None),
        ("Enterprise value", "=B24+B26", "formula", NUM, "b"),
        ("Terminal value as % of EV", "=B26/B27", "formula", PCT, None),
        ("Net debt  (negative = net cash)", nd, "input", NUM, "Net debt at the valuation date (assumptions.json)"),
        ("Non-operating assets", float(A.get("non_operating_assets") or 0), "input", NUM, "Treasury investments and stakes not in FCFF"),
        ("Equity value", "=B27-B29+B30", "formula", NUM, "b"),
        ("Shares outstanding (Cr)", sh_ref, "link", "0.00", None),
        ("Value per share, perpetuity (INR)", "=B31/B32", "formula", INR, "b"),
        ("Current market price (INR)", float(A.get("current_price") or 0), "input", '#,##0.00', cfg.get("price_note", "Price at the cover date")),
        ("Upside / (downside)", "=B33/B34-1", "formula", PCT, None),
        (None, None, None, None, None),
        ("Exit-multiple cross-check and target", None, None, None, None),
        ("Terminal-year EBITDA", f"={LASTC}7", "formula", NUM, None),
        ("Exit EV/EBITDA implied by the perpetuity TV", "=B25/B38", "formula", MULT, None),
        ("Exit multiple (house convention)", float(cfg.get("target_exit_multiple", 10)), "input", MULT, "Target anchor: DCF with an exit multiple"),
        ("Terminal value at exit multiple", "=B40*B38", "formula", NUM, None),
        ("PV of exit terminal value", f"=B41/(1+B22)^{LASTC}17", "formula", NUM, None),
        ("Value per share at exit multiple (INR)", "=(B24+B42-B29+B30)/B32", "formula", INR, "b"),
        ("Rounding step (INR)", float(cfg.get("target_round", 5)), "input", "0", None),
        ("Target price (INR)", "=ROUND(B43/B44,0)*B44", "formula", INR, "b"),
        ("Upside / (downside) to target", "=B45/B34-1", "formula", PCT, None),
        ("BUY above", _rate(cfg.get("buy_above", 15)), "input", PCT, "House rating policy"),
        ("SELL below", -abs(_rate(cfg.get("sell_below", 10))), "input", PCT, "House rating policy"),
        ("Rating", '=IF(B46>B47,"BUY",IF(B46<B48,"SELL","HOLD"))', "formula", "@", "b"),
        (None, None, None, None, None),
        ("Alternative exit multiple", float(A.get("exit_multiple") or 14), "input", MULT, "Cross-check (assumptions.json exit_multiple)"),
        ("Value per share at alternative exit (INR)", f"=(B24+B51*B38/(1+B22)^{LASTC}17-B29+B30)/B32", "formula", INR, None),
    ]
    r = 21
    for lab, v, kind, fmt, note in rows:
        if lab is None:
            r += 1
            continue
        if v is None:
            F.section(r, lab)
        else:
            F.label(r, lab, indent=0 if note == "b" else 1, bold=note == "b",
                    note=None if note in (None, "b") else note, note_col=3)
            F.put(r, 2, v, kind, fmt, bold=note == "b")
        r += 1
    assert ws["A33"].value.strip().startswith("Value per share, perpetuity"), ws["A33"].value
    assert ws["A45"].value.strip().startswith("Target price"), ws["A45"].value
    ws.column_dimensions["A"].width = 50
    for i in range(N):
        ws.column_dimensions[COL(2 + i)].width = 11
    ws.freeze_panes = "B5"
    wb["Scenarios"].cell(live_ref_row, 2, "=DCF!B33")
    wb["Scenarios"].cell(live_ref_row, 2).number_format = INR
    wb["Scenarios"].cell(live_ref_row, 2).font = Font(color=GREEN)

    # ---- Sensitivity ---------------------------------------------------------------------------
    ws = wb.create_sheet("Sensitivity")
    T = Sheet(ws, proto)
    T.title(f"{D['A1'].value.split(' - ')[0]} - Sensitivity",
            "Value per share (INR); every cell recomputes the DCF from the FCFF row")
    pv = (f"SUMPRODUCT(DCF!${FIRSTC}$14:${LASTC}$14,1/(1+{{w}})^DCF!${FIRSTC}$17:${LASTC}$17)")
    T.label(4, "Step (WACC and growth)", bold=True)
    T.put(4, 2, 0.005, "input", PCT)
    T.label(5, "Step (exit multiple)", bold=True)
    T.put(5, 2, 2, "input", MULT)

    def grid(top, title, col_lab, col_val, cell_f):
        T.section(top, title)
        T.put(top + 1, 1, "WACC  ↓  /  " + col_lab + "  →", fmt="@", bold=True)
        for j in range(5):
            T.put(top + 1, 2 + j, col_val(j - 2), fmt=PCT if "growth" in col_lab else MULT, bold=True)
        for i in range(5):
            rr = top + 2 + i
            T.put(rr, 1, f"=DCF!$B$22+({i - 2})*$B$4", fmt='0.00%', bold=True)
            for j in range(5):
                w, x = f"$A{rr}", f"{COL(2 + j)}${top + 1}"
                T.put(rr, 2 + j, cell_f(w, x), fmt=INR, bold=(i == 2 and j == 2))
        return top + 8

    nd_terms = "-DCF!$B$29+DCF!$B$30)/DCF!$B$32"
    nxt = grid(7, "WACC x terminal growth (perpetuity value per share)", "terminal growth",
               lambda k: f"=DCF!$B$23+({k})*$B$4",
               lambda w, g: (f"=IF({g}>={w},NA(),({pv.format(w=w)}+DCF!${LASTC}$14*(1+{g})/({w}-{g})"
                             f"/(1+{w})^DCF!${LASTC}$17{nd_terms})"))
    grid(nxt, "WACC x exit multiple (value per share at exit)", "exit EV/EBITDA",
         lambda k: f"=DCF!$B$40+({k})*$B$5",
         lambda w, m: f"=({pv.format(w=w)}+{m}*DCF!$B$38/(1+{w})^DCF!${LASTC}$17{nd_terms}")
    ws.column_dimensions["A"].width = 30
    for j in range(5):
        ws.column_dimensions[COL(2 + j)].width = 12

    # ---- order, model info ---------------------------------------------------------------------
    names = wb.sheetnames
    want = ["Model info", "Summary", "Scenarios", "Drivers", "Income Statement", "Balance Sheet", "Cash Flow",
            "Schedules", "WACC", "DCF", "Sensitivity", "Checks"]
    want = [n for n in want if n in names] + [n for n in names if n not in want]
    wb._sheets = [wb[n] for n in want]
    wb.active = 0
    info = wb["Model info"]
    r = info.max_row + 2
    info.cell(r, 1, "Valuation").font = Font(bold=True)
    for k, t in enumerate([
            "- Scenarios!B4 selects the case. The switched driver rows on Drivers are green links to it.",
            "- DCF discounts the free cash flow built from these statements; Sensitivity re-prices every cell.",
            "- Blue = input, black = formula, green = link to another sheet. Iterative calculation is on "
            "(interest on average debt is circular).",
            f"- Scenarios '{'/'.join(order)}' were each built by financial-model and valued by model.py; "
            "the Scenarios sheet checks this workbook against those values."]):
        info.cell(r + 1 + k, 1, t)
    wb.calculation.iterate = True
    wb.calculation.iterateCount = 1000
    wb.calculation.iterateDelta = 1e-7
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    return {"varying_rows": [lab for _, lab, _ in varying], "cases": order, "N": N,
            "read": {"value": "DCF!B33", "exit": "DCF!B43", "target": "DCF!B45", "rating": "DCF!B49",
                     "wacc": "DCF!B22", "tv_pct": "DCF!B28", "alt_exit": "DCF!B52", "diff": f"Scenarios!B{live_ref_row + 2}",
                     "weighted": f"Scenarios!D{WTD_ROW}", "fcff1": "DCF!B14", "sens_centre": "Sensitivity!D11"}}


def _lo(args):
    """Run lo_recalc.py: directly where python3 has UNO, through WSL on Windows."""
    import subprocess
    here = os.path.join(os.path.dirname(os.path.abspath(__file__)), "lo_recalc.py")
    if os.name == "nt":
        wsl = lambda p: "/mnt/" + p[0].lower() + p[2:].replace("\\", "/")
        conv = [wsl(a) if (len(a) > 2 and a[1] == ":") else a for a in args]
        cmd = ["wsl", "-d", os.environ.get("ERIP_WSL_DISTRO", "Ubuntu"), "--", "python3", wsl(here)] + conv
    else:
        cmd = ["python3", here] + args
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        raise SystemExit(f"LibreOffice recalc failed:\n{r.stdout[-800:]}\n{r.stderr[-800:]}")
    return json.loads(r.stdout.strip().splitlines()[-1])


def recalc_and_verify(raw, out, info):
    """Switch through every case in LibreOffice; each must reproduce its recorded value with no
    formula errors. Then save the workbook on the base case, values cached beside the formulas."""
    rd = info["read"]
    report = {}
    for i, case in enumerate(info["cases"], 1):
        res = _lo([raw, "--set", f"Scenarios!B4={i}", "--read", rd["value"], "--read", rd["target"],
                   "--read", rd["rating"], "--read", rd["diff"], "--errors"])
        v = res["read"]
        report[case] = {"value": round(v[rd["value"]], 2), "target": v[rd["target"]], "rating": v[rd["rating"]],
                        "diff_vs_python": v[rd["diff"]], "errors": res["errors"], "passes": res["passes"]}
        if res["errors"] or v[rd["diff"]] != 0:
            raise SystemExit(f"{case}: workbook does not reproduce the Python value: {report[case]}")
    _lo([raw, "--set", "Scenarios!B4=1", "--out", out])
    return report


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--root", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--recalc", action="store_true",
                    help="verify every scenario in LibreOffice and save with cached values (needs LibreOffice + UNO)")
    a = ap.parse_args()
    out = os.path.abspath(a.out)
    raw = out[:-5] + ".formulas.xlsx" if a.recalc else out
    info = build(os.path.abspath(a.root), raw)
    print("switched driver rows:", ", ".join(info["varying_rows"]))
    if a.recalc:
        for case, r in recalc_and_verify(raw, out, info).items():
            print(f"  {case:5} INR {r['value']:>8}  target {r['target']:>6}  {r['rating']:4}  "
                  f"diff vs Python {r['diff_vs_python']}  errors {len(r['errors'])}  passes {r['passes']}")
        os.remove(raw)
    print("->", out)
