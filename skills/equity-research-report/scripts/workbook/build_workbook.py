#!/usr/bin/env python3
"""build_workbook.py - produce a professional Excel workbook from model.json.

 python build_workbook.py --model data/three_statement/model.json \
 --assumptions data/assumptions.json --outdir data/three_statement/

Writes: <Company>_ERIP_Model.xlsx

Companion to financial-model. Consumes model.json (authoritative solved
state) and assumptions.json. Does NOT recalculate the model.

Requires: openpyxl
"""
from __future__ import annotations

import argparse, json, os, sys
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl required: pip install openpyxl")

NAVY = "FF1F3864"
HDR_FONT = "FFFFFFFF"
GREY = "FF808080"
BLACK = "FF000000"
hdr_fill = PatternFill("solid", fgColor=NAVY)
sec_fill = PatternFill("solid", fgColor="FFEDEDED")
fcst_fill = PatternFill("solid", fgColor="FFFFF7E6")
inp_fill = PatternFill("solid", fgColor="FFFFFDE7")
red_fill = PatternFill("solid", fgColor="FFFFC7CE")

DATA_COL0 = 2
FIRST_ROW = 5

def col_letter(n):
    return get_column_letter(n)

def cell_hdr(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font = Font(bold=True, color=HDR_FONT, size=10)
    c.fill = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c

def cell_val(ws, row, col, value, fmt=None, bold=False, color=BLACK,
    fill=None, indent=0):
    c = ws.cell(row, col, value)
    c.font = Font(bold=bold, size=10, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    c.alignment = Alignment(indent=indent, vertical="top", wrap_text=True)
    return c

def write_info(ws, model, checks, recon, notes):
    ws.sheet_view.showGridLines = False
    ws["A1"] = model.get("company", "") + " - Three-Statement Model"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    ws["A2"] = "Currency: " + str(model.get("currency", "")) + " | Built: " + str(model.get("built", "")) + " | Sector: " + str(model.get("sector", "N/A"))
    ws["A2"].font = Font(size=10, color=GREY, italic=True)
    r = 4

    def section(title):
        nonlocal r
        ws.cell(r, 1, title).font = Font(bold=True, size=12, color=NAVY)
        r += 1

    def kv(label, value, bold=False):
        nonlocal r
        ws.cell(r, 1, label).font = Font(bold=bold, size=10)
        ws.cell(r, 2, value).font = Font(bold=bold, size=10)
        r += 1

    section("Model Parameters")
    kv("Company", model.get("company", ""), bold=True)
    kv("Currency", model.get("currency", ""))
    kv("Sector", model.get("sector", "—"))
    periods = model.get("periods", {})
    hist = periods.get("hist", [])
    fcst = periods.get("fcst", [])
    h_str = hist[0]+" - "+hist[-1]+" ("+str(len(hist))+" yr)" if hist else "—"
    kv("History", h_str)
    f_str = fcst[0]+" - "+fcst[-1]+" ("+str(len(fcst))+" yr)" if fcst else "—"
    kv("Forecast", f_str)
    kv("Depreciation basis", model.get("options",{}).get("dep_basis","net_block"))
    solve = model.get("solve", {})
    kv("Circular solve", str(solve.get("iterations","?"))+" passes, residual "+str(solve.get("residual","?")))
    kv("Checks status", "ALL PASS" if checks.get("pass") else "FAILED", bold=not checks.get("pass"))
    r += 1
    section("Integrity Checks")
    for it in checks.get("items", []):
        ok = it.get("ok", False)
        status = "OK" if ok else "FAIL"
        color = BLACK if ok else "FFC00000"
        ws.cell(r, 1, status+" "+it.get("label","")).font = Font(size=10, color=color)
        ws.cell(r, 2, it.get("severity", ""))
        ws.cell(r, 3, it.get("worst", ""))
        ws.cell(r, 4, it.get("period", ""))
        r += 1
    if recon:
        r += 1
        section("Screener vs Annual Report Reconciliation")
        for h, t in enumerate(["Line","Period","Screener","Annual Report","Diff %"]):
            ws.cell(r, 1+h, t).font = Font(bold=True, size=9)
            r += 1
        for d in recon:
            ws.cell(r, 1, d.get("row","")).font = Font(size=9)
            ws.cell(r, 2, d.get("period","")).font = Font(size=9)
            ws.cell(r, 3, d.get("screener","")).font = Font(size=9)
            ws.cell(r, 4, d.get("annual_report","")).font = Font(size=9)
            ws.cell(r, 5, d.get("diff_pct","")).font = Font(size=9)
            r += 1
    citations = model.get("citations", {})
    if citations:
        r += 1
        section("Annual-Report Citations")
        for k, c in citations.items():
            src = c.get("source","") if isinstance(c, dict) else str(c)
            page = c.get("page","") if isinstance(c, dict) else ""
            ws.cell(r, 1, k).font = Font(size=10)
            ws.cell(r, 2, src+" p."+str(page)).font = Font(size=10)
            r += 1
    if notes:
        r += 1
        section("Build Notes")
        for n in notes:
            ws.cell(r, 1, "- "+n).font = Font(size=10)
            r += 1

    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

def write_data_sheet(ws, title, model, sectioned_rows, fmt_map, n_hist, ncols):
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, title)
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    cols = model["periods"]["all"]
    ws.cell(3, 1, "Actuals")
    ws.cell(3, 1).font = Font(bold=True, color=HDR_FONT, size=10)
    ws.cell(3, 1).fill = hdr_fill
    ws.cell(3, 1).alignment = Alignment(horizontal="center")
    for j, p in enumerate(cols):
        c = ws.cell(3, DATA_COL0 + j, p)
        c.font = Font(bold=True, color=HDR_FONT, size=10)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    rows_data = model.get("rows", {})
    row_num = FIRST_ROW
    no_gm = any("no gross margin" in str(n_).lower() for n_ in model.get("notes", []))
    # Where gross margin was set equal to EBITDA margin, every operating cost sits in COGS and opex is zero:
    # keep the cost line (relabelled), hide gross profit, gross margin and the empty opex line.
    GM_KEYS = {"gross_profit", "gross_margin", "opex"}
    for item in sectioned_rows:
        key, label, indent, kind = item[:4]
        spec = item[4] if len(item) > 4 else {}
        if no_gm and (key in GM_KEYS or (key is None and "gross" in label.lower())):
            continue
        if no_gm and key == "cogs":
            label = "Operating costs"
        if key is None:
            ws.cell(row_num, 1, (" " * indent) + label).font = Font(bold=True, size=10, color=NAVY)
            ws.cell(row_num, 1).fill = sec_fill
            for j in range(ncols + 1):
                ws.cell(row_num, 1 + j).fill = sec_fill
            row_num += 1
            continue
        is_total = bool(spec.get("total"))
        ws.cell(row_num, 1, label)
        ws.cell(row_num, 1).font = Font(bold=is_total, size=10)
        ws.cell(row_num, 1).alignment = Alignment(indent=indent)
        row_data = rows_data.get(key, {})
        values = spec.get("values", row_data.get("values", []))
        blank = spec.get("blank", set())
        fmt = spec.get("fmt", fmt_map.get(key, "#,##0.00"))
        is_input = kind == "input"
        is_chk = kind == "chk"
        for j in range(ncols):
            v = values[j] if j < len(values) else None
            # A rate with no prior period (first-year growth) arrives as -1.0: show it blank, not -100%.
            if (key in ("rev_growth",) and j == 0) or j in blank:
                v = None
            fill = None
            if j >= n_hist:
                fill = fcst_fill
            if is_input and j >= n_hist:
                fill = inp_fill
            if is_chk:
                fill = red_fill
            c = ws.cell(row_num, DATA_COL0 + j, None if v is None else round(float(v), 6))
            c.number_format = fmt
            c.font = Font(size=10, bold=is_total, color="FF0000FF" if (is_input and j >= n_hist) else BLACK)
            if is_total:
                c.border = Border(top=Side(style="thin", color="FF7F7F7F"))
            if fill:
                c.fill = fill
        row_num += 1      # one row per line item: this sat inside the period loop and staircased every sheet
    ws.column_dimensions["A"].width = 44
    for j in range(ncols):
        ws.column_dimensions[col_letter(DATA_COL0 + j)].width = 14
    ws.freeze_panes = ws.cell(FIRST_ROW, DATA_COL0)

def write_checks(ws, model, n_hist):
    """One row per integrity check: status first, then what the number is. A value of 0 on an identity
    check means it ties; for a floor check (net block, equity) the worst value is the lowest seen."""
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Integrity Checks")
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    items = model.get("checks", {}).get("items", [])
    n_block = sum(1 for it in items if not it.get("ok", True) and it.get("severity") != "advisory")
    ws.cell(2, 1, f"{len(items)} checks: {n_block} blocking failure(s); advisory flags are explained, not errors."
            ).font = Font(italic=True, size=10)
    hdr = ["Check", "Status", "Severity", "Worst value", "Period", "What it tests"]
    for j, h in enumerate(hdr):
        c = ws.cell(4, 1 + j, h)
        c.font = Font(bold=True, color=HDR_FONT, size=10)
        c.fill = hdr_fill
    r = 5
    for it in items:
        ok, sev = it.get("ok", True), it.get("severity", "blocking")
        status = "PASS" if ok else ("ADVISORY" if sev == "advisory" else "FAIL")
        ws.cell(r, 1, it.get("label", it.get("check"))).font = Font(size=10)
        c = ws.cell(r, 2, status)
        c.font = Font(bold=True, size=10, color="FF008000" if ok else ("FFB26B00" if sev == "advisory" else "FFC00000"))
        ws.cell(r, 3, sev).font = Font(size=10)
        v = it.get("worst")
        c = ws.cell(r, 4, v if isinstance(v, (int, float)) else None)
        c.number_format = "#,##0.00"
        ws.cell(r, 5, it.get("period")).font = Font(size=10)
        c = ws.cell(r, 6, it.get("detail", ""))
        c.font = Font(size=9, italic=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    for col, wd in zip("ABCDEF", (52, 11, 11, 13, 10, 70)):
        ws.column_dimensions[col].width = wd
    ws.freeze_panes = "A5"

def write_summary(ws, model, n_hist):
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Summary - Key Metrics")
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    rows_data = model.get("rows", {})
    cols = model["periods"]["all"]
    ncols = len(cols)
    ws.cell(3, 1, "Actuals")
    ws.cell(3, 1).font = Font(bold=True, color=HDR_FONT, size=10)
    ws.cell(3, 1).fill = hdr_fill
    ws.cell(3, 1).alignment = Alignment(horizontal="center")
    for j, p in enumerate(cols):
        c = ws.cell(3, DATA_COL0 + j, p)
        c.font = Font(bold=True, color=HDR_FONT, size=10)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    summary_items = [
        ("revenue", "Revenue", "#,##0.00"),
        ("ebitda", "EBITDA", "#,##0.00"),
        ("ebitda_margin", "EBITDA margin", "0.0%"),
        ("ebit", "EBIT", "#,##0.00"),
        ("pat", "Profit after tax", "#,##0.00"),
        ("_s_eps", "EPS (INR)", "#,##0.00"),
        ("_s_fcff", "FCFF (history: reported CFO less capex)", "#,##0.00"),
        ("_s_nd", "Net debt", "#,##0.00"),
        ("debt_total", "Total debt", "#,##0.00"),
        ("_s_cash", "Cash and bank", "#,##0.00"),
    ]
    row_num = FIRST_ROW
    for key, label, fmt in summary_items:
        ws.cell(row_num, 1, label)
        ws.cell(row_num, 1).font = Font(bold=True, size=10)
        row_data = rows_data.get(key, {})
        values = row_data.get("values", [])
        for j in range(ncols):
            v = values[j] if j < len(values) else None
            c = ws.cell(row_num, DATA_COL0 + j, None if v is None else round(float(v), 6))
            c.number_format = fmt
            if j >= n_hist:
                c.fill = fcst_fill
        row_num += 1
    ws.column_dimensions["A"].width = 30
    for j in range(ncols):
        ws.column_dimensions[col_letter(DATA_COL0 + j)].width = 14
    ws.freeze_panes = ws.cell(FIRST_ROW, DATA_COL0)

def _frac(v):
    """assumptions.json carries rates as percent (12.0) or decimal (0.12); return a fraction."""
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return None
    return fv / 100 if abs(fv) > 1 else fv


SC_COLOR = {"base": NAVY, "bull": "FF008000", "bear": "FFC00000"}


def write_scenarios(ws, model, assumptions, n_hist, dcf_model=None):
    """One row per scenario per driver. A scenario that does not state a driver inherits the base
    value from the top level of assumptions.json, which is how model.py runs it."""
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Scenario Comparison")
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    scenarios = {k: v for k, v in (assumptions.get("scenarios") or {}).items() if not k.startswith("_")}
    if not scenarios:
        ws.cell(3, 1, "No scenarios defined in assumptions.json")
        ws.column_dimensions["A"].width = 40
        return
    order = [k for k in ("bull", "base", "bear") if k in scenarios] + [k for k in scenarios if k not in ("bull", "base", "bear")]
    fcst = model["periods"]["fcst"]
    ncols = len(fcst)
    ws.cell(3, 1, "Driver / case")
    for j, p in enumerate(["Case"] + fcst):
        c = ws.cell(3, 2 + j, p)
        c.font = Font(bold=True, color=HDR_FONT, size=10)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    ws.cell(3, 1).font = Font(bold=True, color=HDR_FONT, size=10)
    ws.cell(3, 1).fill = hdr_fill
    metrics = [("revenue_growth", "Revenue growth"), ("ebitda_margin", "EBITDA margin"),
               ("ebit_margin", "EBIT margin"), ("capex_pct_sales", "Capex / sales"),
               ("terminal_growth", "Terminal growth")]
    row_num = 4
    for key, label in metrics:
        base_val = assumptions.get(key)
        if base_val is None and not any(key in scenarios[k] for k in order):
            continue
        ws.cell(row_num, 1, label).font = Font(bold=True, size=10)
        for sc in order:
            vals = scenarios[sc].get(key, base_val)
            if vals is None:
                continue
            if not isinstance(vals, list):
                vals = [vals] * ncols
            ws.cell(row_num, 2, sc).font = Font(bold=True, size=10, color=SC_COLOR.get(sc, NAVY))
            for j in range(ncols):
                v = vals[j] if j < len(vals) else vals[-1]
                fv = _frac(v)
                c = ws.cell(row_num, 3 + j, fv)
                c.number_format = "0.0%"
                c.font = Font(size=10, color=SC_COLOR.get(sc, NAVY))
            row_num += 1
        row_num += 1
    ws.cell(row_num, 1, "Probability").font = Font(bold=True, size=10)
    cases = (dcf_model or {}).get("scenarios", {}).get("cases", {})
    if cases:
        ws.cell(row_num, 3, "INR / share").font = Font(bold=True, size=10)
        ws.cell(row_num, 4, "TV % of EV").font = Font(bold=True, size=10)
    row_num += 1
    for sc in order:
        ws.cell(row_num, 1, sc).font = Font(bold=True, size=10, color=SC_COLOR.get(sc, NAVY))
        c = ws.cell(row_num, 2, scenarios[sc].get("probability"))
        c.number_format = "0%"
        if sc in cases:
            ws.cell(row_num, 3, cases[sc].get("intrinsic_value_per_share"))
            ws.cell(row_num, 4, (cases[sc].get("TV_as_%_of_EV") or 0) / 100).number_format = "0%"
        row_num += 1
    row_num += 1
    ws.cell(row_num, 1, "Rationale").font = Font(bold=True, size=11, color=NAVY)
    row_num += 1
    for sc in order:
        why = scenarios[sc].get("_why", "")
        if why:
            ws.cell(row_num, 1, sc + ": " + why).font = Font(size=10, italic=True)
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2 + ncols)
            ws.cell(row_num, 1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[row_num].height = 60
            row_num += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    for j in range(ncols):
        ws.column_dimensions[col_letter(3 + j)].width = 11


def write_valuation(ws, model, assumptions, dcf_model=None):
    """WACC build, the three-statement FCFF bridge laid out by period, and, when model.py's output is
    supplied with --dcf, the equity bridge to a per-share value. Never recalculates: it lays out
    what the authoritative files already hold."""
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "DCF Valuation Summary")
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    row_num = 3

    def section(title):
        nonlocal row_num
        ws.cell(row_num, 1, title).font = Font(bold=True, size=12, color=NAVY)
        row_num += 1

    def kv(label, value, fmt=None):
        nonlocal row_num
        ws.cell(row_num, 1, label).font = Font(bold=True, size=10)
        c = ws.cell(row_num, 2, value if isinstance(value, (int, float)) else str(value))
        c.font = Font(size=10)
        if fmt and isinstance(value, (int, float)):
            c.number_format = fmt
        row_num += 1

    dm = dcf_model or {}
    wbld = (dm.get("dcf") or {}).get("wacc_build")
    section("WACC build" + (" (model.py)" if wbld else " (assumptions.json)"))
    if wbld:
        for k, v in wbld.items():
            if isinstance(v, dict):
                for kk, vv in v.items():
                    kv("  " + kk, vv, "0.000")
            elif isinstance(v, (int, float)):
                kv(k, v / 100 if k not in ("Beta",) and abs(v) > 1 else v, "0.00" if k == "Beta" else "0.00%")
            else:
                kv(k, v)
    else:
        for k in ("rf", "erp", "beta", "cost_of_debt", "tax_rate", "terminal_growth", "target_debt_weight"):
            if k in assumptions:
                v = assumptions[k]
                kv(k, v if k == "beta" else _frac(v), "0.00" if k == "beta" else "0.00%")
    row_num += 1

    bridge = model.get("bridge", {}) or {}
    periods = bridge.get("periods") or model["periods"]["fcst"]
    series = [(k, v) for k, v in bridge.items() if isinstance(v, list) and k != "periods"]
    if series:
        section("Three-statement FCFF bridge (INR Cr)")
        for j, p in enumerate(["Line"] + list(periods)):
            c = ws.cell(row_num, 1 + j, p)
            c.font = Font(bold=True, color=HDR_FONT, size=10)
            c.fill = hdr_fill
        row_num += 1
        for k, vals in series:
            ws.cell(row_num, 1, k).font = Font(bold=(k == "fcff"), size=10)
            for j, v in enumerate(vals):
                c = ws.cell(row_num, 2 + j, v)
                c.number_format = "0.0%" if k == "tax_rate" else "#,##0"
            row_num += 1
        row_num += 1

    br = (dm.get("dcf") or {}).get("bridge")
    if br:
        section("Enterprise value to equity value (model.py)")
        for k, v in br.items():
            kv(k, v, "#,##0.00" if isinstance(v, float) and abs(v) < 1000 else "#,##0")
        tcc = (dm.get("dcf") or {}).get("terminal_cross_check") or {}
        if tcc:
            row_num += 1
            section("Terminal value cross-check")
            for k, v in tcc.items():
                kv(k, v, "#,##0.0" if isinstance(v, (int, float)) else None)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14
    for j in range(len(periods)):
        ws.column_dimensions[col_letter(3 + j)].width = 11


def write_sensitivity(ws, assumptions, dcf_model=None):
    """WACC x terminal-growth grid of equity value per share, from model.py's output (--dcf)."""
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, "Sensitivity - Equity Value per Share (INR)")
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    sens = ((dcf_model or {}).get("dcf") or {}).get("sensitivity") or assumptions.get("dcf", {}).get("sensitivity", {})
    wacc_vals, tg_vals, grid = sens.get("wacc", []), sens.get("terminal_growth", []), sens.get("grid", [])
    if not (wacc_vals and tg_vals and grid):
        ws.cell(3, 1, "No sensitivity grid: pass --dcf data/model.json (model.py output)")
        ws.column_dimensions["A"].width = 60
        return
    br = ((dcf_model or {}).get("dcf") or {}).get("bridge", {})
    base, cmp_ = br.get("Intrinsic value per share (INR)"), br.get("Current market price (INR)")
    ws.cell(2, 1, f"Base case INR {base} vs market price INR {cmp_}" if base else "").font = Font(italic=True, size=10)
    c = ws.cell(4, 1, "WACC x g")
    c.font = Font(bold=True, color=HDR_FONT, size=10)
    c.fill = hdr_fill
    for j, tg in enumerate(tg_vals):
        c = ws.cell(4, 2 + j, tg)
        c.font = Font(bold=True, size=10, color=HDR_FONT)
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    for i, wv in enumerate(wacc_vals):
        ws.cell(5 + i, 1, wv).font = Font(bold=True, size=10)
        for j, v in enumerate(grid[i] if i < len(grid) else []):
            c = ws.cell(5 + i, 2 + j, v)
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="center")
            c.font = Font(size=10, bold=(v == base))
    ws.column_dimensions["A"].width = 14
    for j in range(len(tg_vals)):
        ws.column_dimensions[col_letter(2 + j)].width = 12


def build_workbook(model_path, assumptions_path, out_dir, filename=None, dcf_path=None):
    dcf_model = json.load(open(dcf_path, encoding="utf-8")) if dcf_path else None
    with open(model_path, encoding="utf-8") as fh:
        model = json.load(fh)
    assumptions = {}
    if assumptions_path and Path(assumptions_path).exists():
        with open(assumptions_path, encoding="utf-8") as fh:
            assumptions = json.load(fh)
    company = model.get("company", "Company")
    safe_name = company.replace(" ", "_").replace("/", "_").replace("&", "")
    if not filename:
        filename = safe_name + "_ERIP_Model.xlsx"
    out_path = os.path.join(out_dir, filename)
    periods = model.get("periods", {})
    hist = periods.get("hist", [])
    fcst = periods.get("fcst", [])
    n_hist = len(hist)
    ncols = len(hist) + len(fcst)
    wb = Workbook()
    wb.remove(wb.active)
    for name in ["Model info", "Drivers", "Income Statement", "Balance Sheet",
        "Cash Flow", "Schedules", "Checks", "Summary",
        "Scenarios", "Valuation", "Sensitivity"]:
        wb.create_sheet(name)
    checks = model.get("checks", {"pass": True, "items": []})
    recon = model.get("reconciliation", [])
    notes = model.get("notes", [])
    rows_data = model.get("rows", {})
    write_info(wb["Model info"], model, checks, recon, notes)

    # ---- presentation helpers: what a reader should see, without changing the model's numbers
    V = lambda k: list(rows_data.get(k, {}).get("values", [None] * ncols)) + [None] * ncols
    H = set(range(n_hist))
    no_gm = any("no gross margin" in str(n_).lower() for n_ in notes)
    hist_zero = lambda k: {j for j in range(n_hist) if not V(k)[j]}
    DRV_SPEC = {
        "rev_growth": {"blank": {0}}, "debt_growth": {"blank": {0}},
        "capex_pct_sales": {"blank": {0}}, "dep_pct_nb": {"blank": {0}},
        "interest_income_rate": {"blank": H}, "term_debt_repay": {"blank": H},
        "equity_issued": {"blank": H}, "min_cash": {"blank": H},
        "dpo": {"blank": hist_zero("dpo")},
    }

    # Drivers sheet
    driver_keys = [
        ("rev_growth", "Revenue growth", 0, "input"),
        ("ebitda_margin", "EBITDA margin", 0, "input"),
        ("gross_margin", "Gross margin", 0, "input"),
        ("other_income_pct", "Other income (% sales)", 0, "input"),
        ("capex_pct_sales", "Capex (% sales)", 0, "input"),
        ("dep_pct_nb", "Depreciation (% opening NB)", 0, "input"),
        ("dep_pct_sales", "Depreciation (% sales)", 0, "input"),
        ("cwip_pct_sales", "CWIP (% sales)", 0, "input"),
        ("dso", "Receivable days", 0, "input"),
        ("dio", "Inventory days", 0, "input"),
        ("dpo", "Payable days", 0, "input"),
        ("tax_rate", "Effective tax rate", 0, "input"),
        ("cost_of_debt", "Interest rate on debt", 0, "input"),
        ("interest_income_rate", "Yield on cash", 0, "input"),
        ("debt_growth", "Term debt growth", 0, "input"),
        ("term_debt_repay", "Scheduled debt repayment", 0, "input"),
        ("equity_issued", "Equity issued", 0, "input"),
        ("payout_ratio", "Dividend payout", 0, "input"),
        ("min_cash", "Minimum cash balance", 0, "input"),
        ("other_ca_pct", "Other CA (% sales)", 0, "input"),
        ("other_liab_pct", "Other liabilities (% sales)", 0, "input"),
    ]
    driver_fmts = {
        "rev_growth": "0.00%", "ebitda_margin": "0.00%", "gross_margin": "0.00%",
        "other_income_pct": "0.00%", "capex_pct_sales": "0.00%",
        "dep_pct_nb": "0.00%", "dep_pct_sales": "0.00%", "cwip_pct_sales": "0.00%",
        "dso": "0.0", "dio": "0.0", "dpo": "0.0",
        "tax_rate": "0.00%", "cost_of_debt": "0.00%", "interest_income_rate": "0.00%",
        "debt_growth": "0.00%", "term_debt_repay": "#,##0.00", "equity_issued": "#,##0.00",
        "payout_ratio": "0.00%", "min_cash": "#,##0.00",
        "other_ca_pct": "0.00%", "other_liab_pct": "0.00%",
    }
    driver_keys = [(k, l, i, kd, DRV_SPEC.get(k, {})) for k, l, i, kd in driver_keys]
    write_data_sheet(wb["Drivers"], company + " - Drivers",
        model, driver_keys, driver_fmts, n_hist, ncols)

    # Income Statement
    _rev, _eb = V("revenue"), V("ebitda")
    _opc = [(r - e) if None not in (r, e) else None for r, e in zip(_rev[:ncols], _eb[:ncols])]
    _ebm = [(e / r) if (r and e is not None) else None for r, e in zip(_rev[:ncols], _eb[:ncols])]
    is_items = [(None, "INCOME STATEMENT", 0, "head"), ("revenue", "Revenue", 0, "actual")]
    if no_gm:
        is_items.append(("_opcost", "Operating costs (revenue less EBITDA)", 0, "calc", {"values": _opc}))
    else:
        is_items += [("cogs", "Cost of goods sold", 0, "actual"),
                     ("gross_profit", "Gross profit", 0, "calc", {"total": True}),
                     ("opex", "Operating expenses", 0, "calc")]
    is_items += [
        ("ebitda", "EBITDA", 0, "actual", {"total": True}),
        ("_ebm", "EBITDA margin", 1, "calc", {"values": _ebm, "fmt": "0.0%"}),
        ("dep_charge", "Depreciation and amortisation", 0, "actual"),
        ("ebit", "EBIT", 0, "calc", {"total": True}),
        ("other_income", "Other income", 0, "actual"),
        ("interest_expense", "Interest expense", 0, "actual"),
    ]
    if any(V("exceptional")[:ncols]):
        is_items.append(("exceptional", "Exceptional items", 0, "actual"))
    is_items += [
        ("pbt", "Profit before tax", 0, "actual", {"total": True}),
        ("tax", "Tax", 0, "actual"),
        ("pat", "Profit after tax", 0, "actual", {"total": True}),
        ("eps", "Earnings per share (INR)", 0, "calc", {"blank": {j for j in range(n_hist) if abs(V("eps")[j] or 0) < 0.01}}),
    ]
    is_fmts = {it[0]: "#,##0.00" for it in is_items if it[0]}
    write_data_sheet(wb["Income Statement"], company + " - Income Statement",
        model, is_items, is_fmts, n_hist, ncols)

    # Balance Sheet
    ca = ["cash", "receivables", "inventory", "other_ca"]
    nca = ["net_block", "cwip", "investments"]
    eq = ["share_capital", "reserves"]
    debt_k = ["term_debt", "revolver"]
    lia = ["payables", "other_liab"]
    def _col_sums(keys):
        return [sum(float(rows_data.get(k, {}).get("values", [0]*ncols)[j] or 0) for k in keys) for j in range(ncols)]
    extended = dict(rows_data)
    extended["total_current_assets"] = {"label":"Total current assets","sheet":"BS","kind":"calc","values":_col_sums(ca)}
    extended["total_non_current_assets"] = {"label":"Total non-current assets","sheet":"BS","kind":"calc","values":_col_sums(nca)}
    extended["total_equity"] = {"label":"Total equity","sheet":"BS","kind":"calc","values":_col_sums(eq)}
    extended["total_debt"] = {"label":"Total debt","sheet":"BS","kind":"calc","values":_col_sums(debt_k)}
    extended["total_liabilities"] = {"label":"Total liabilities","sheet":"BS","kind":"calc","values":_col_sums(lia)}
    old_rows = model["rows"]
    model["rows"] = extended
    bs_items = [
        (None, "ASSETS", 0, "head"),
        ("cash", "Cash and bank", 0, "actual"),
        ("receivables", "Trade receivables", 0, "actual"),
        ("inventory", "Inventories", 0, "actual"),
        ("other_ca", "Other current assets", 0, "actual"),
        ("total_current_assets", "Total current assets", 1, "calc"),
        (None, "Non-current assets", 1, "head"),
        ("net_block", "Net fixed assets", 0, "actual"),
        ("cwip", "Capital work in progress", 0, "actual"),
        ("investments", "Investments", 0, "actual"),
        ("total_non_current_assets", "Total non-current assets", 1, "calc"),
        ("total_assets", "Total assets", 0, "calc"),
        (None, "EQUITY AND LIABILITIES", 0, "head"),
        ("share_capital", "Share capital", 0, "actual"),
        ("reserves", "Reserves and surplus", 0, "actual"),
        ("total_equity", "Total equity", 1, "calc"),
        ("term_debt", "Term borrowings", 0, "actual"),
        ("revolver", "Revolver / short-term plug", 0, "actual"),
        ("total_debt", "Total debt", 1, "calc"),
        ("payables", "Trade payables", 0, "actual", {"blank": hist_zero("payables")}),
        ("other_liab", "Other liabilities", 0, "actual"),
        ("total_liabilities", "Total liabilities", 1, "calc"),
        ("total_liab_eq", "Total liabilities and equity", 0, "calc"),
    ]
    bs_fmts = {it[0]: "#,##0.00" for it in bs_items if it[0]}
    write_data_sheet(wb["Balance Sheet"], company + " - Balance Sheet",
        model, bs_items, bs_fmts, n_hist, ncols)
    model["rows"] = old_rows

    # Cash Flow
    _rep = lambda mk, rk: [(V(rk)[j] if j < n_hist else V(mk)[j]) for j in range(ncols)]
    _cfo, _cfi, _cff = _rep("cfo", "rep_cfo"), _rep("cfi", "rep_cfi"), _rep("cff", "rep_cff")
    _net = [(None if None in (a_, b_, c_) else a_ + b_ + c_) if j < n_hist else V("net_change_cash")[j]
            for j, (a_, b_, c_) in enumerate(zip(_cfo, _cfi, _cff))]
    FH = {"blank": H}
    cf_items = [
        (None, "CASH FROM OPERATING ACTIVITIES", 0, "head"),
        ("cf_pat", "Profit after tax", 1, "calc", FH),
        ("cf_dep", "Depreciation and amortisation", 1, "calc", FH),
        ("cf_wc", "(Increase)/decrease in net operating assets", 1, "calc", FH),
        ("cfo", "Cash from operations", 0, "calc", {"values": _cfo, "total": True}),
        (None, "CASH FROM INVESTING ACTIVITIES", 0, "head"),
        ("cf_capex", "Capital expenditure", 1, "calc", FH),
        ("cf_invest", "Investments", 1, "calc", FH),
        ("cfi", "Cash from investing", 0, "calc", {"values": _cfi, "total": True}),
        (None, "CASH FROM FINANCING ACTIVITIES", 0, "head"),
        ("cf_debt", "Term debt drawn/(repaid)", 1, "calc", FH),
        ("cf_revolver", "Revolver drawn/(repaid)", 1, "calc", FH),
        ("cf_div", "Dividends paid", 1, "calc", FH),
        ("cf_equity", "Equity issued", 1, "calc", FH),
        ("cff", "Cash from financing", 0, "calc", {"values": _cff, "total": True}),
        (None, "NET CHANGE IN CASH", 0, "head"),
        ("net_change_cash", "Net change in cash", 0, "calc", {"values": _net, "total": True}),
        ("cash_open", "Opening cash", 1, "calc", FH),
        ("cash_close", "Closing cash", 1, "calc", FH),
    ]
    cf_fmts = {it[0]: "#,##0.00" for it in cf_items if it[0]}
    write_data_sheet(wb["Cash Flow"], company + " - Cash Flow (history: reported totals; line detail is modelled for forecast years)",
        model, cf_items, cf_fmts, n_hist, ncols)

    # Schedules
    _nb = V("net_block")
    _nbo = [(_nb[j - 1] if j else None) if j < n_hist else V("nb_open")[j] for j in range(ncols)]
    _nbc = [_nb[j] if j < n_hist else V("nb_close")[j] for j in range(ncols)]
    sched_items = [
        ("nb_open", "Opening net block", 0, "calc", {"values": _nbo}),
        ("capex", "Capital expenditure", 0, "actual", {"blank": {0}}),
        ("sch_dep", "Depreciation charge", 0, "calc"),
        ("nb_close", "Closing net block", 0, "calc", {"values": _nbc, "total": True}),
        (None, "Net operating assets", 0, "head"),
        ("nwc", "Net operating assets", 0, "calc"),
        ("d_nwc", "Change in net operating assets", 0, "calc", {"blank": {0}}),
        ("wc_days", "Net operating assets (days)", 0, "calc"),
    ]
    sched_fmts = {"nb_open": "#,##0.00", "capex": "#,##0.00",
        "sch_dep": "#,##0.00", "nb_close": "#,##0.00",
        "nwc": "#,##0.00", "d_nwc": "#,##0.00", "wc_days": "0.0"}
    write_data_sheet(wb["Schedules"], company + " - Schedules",
        model, sched_items, sched_fmts, n_hist, ncols)

    write_checks(wb["Checks"], model, n_hist)
    # Summary history from reported figures; the model's reconstructed history has no opening balances
    _cash = V("cash")
    _dbt = V("debt_total")
    _capx = V("capex")
    _rs = model["rows"]
    _rs["_s_cash"] = {"values": _cash[:ncols]}
    _rs["_s_nd"] = {"values": [(d - c) if None not in (d, c) else None for d, c in zip(_dbt[:ncols], _cash[:ncols])]}
    _rs["_s_fcff"] = {"values": [((V("rep_cfo")[j] or 0) - (_capx[j] or 0)) if (j < n_hist and j > 0) else
                                 (None if j < n_hist else V("s_fcff")[j]) for j in range(ncols)]}
    _rs["_s_eps"] = {"values": [(None if (j < n_hist and abs(V("eps")[j] or 0) < 0.01) else V("eps")[j]) for j in range(ncols)]}
    write_summary(wb["Summary"], model, n_hist)
    write_scenarios(wb["Scenarios"], model, assumptions, n_hist, dcf_model)
    write_valuation(wb["Valuation"], model, assumptions, dcf_model)
    write_sensitivity(wb["Sensitivity"], assumptions, dcf_model)

    try:
        wb.calculation.iterate = True
        wb.calculation.iterateCount = 200
        wb.calculation.iterateDelta = 1e-6
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        pass

    os.makedirs(out_dir, exist_ok=True)
    wb.save(out_path)
    return {
        "path": out_path, "filename": filename,
        "sheets": wb.sheetnames, "n_hist": n_hist, "n_fcst": len(fcst),
        "checks_pass": checks.get("pass", True), "n_checks": len(checks.get("items", [])),
    }

def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--model", required=True, help="Path to model.json")
    ap.add_argument("--assumptions", default=None, help="Path to assumptions.json")
    ap.add_argument("--outdir", required=True, help="Output directory")
    ap.add_argument("--filename", default=None, help="Output filename")
    ap.add_argument("--dcf", default=None, help="model.py output (data/model.json): adds the equity bridge and scenario values")
    args = ap.parse_args()
    result = build_workbook(args.model, args.assumptions, args.outdir, args.filename, args.dcf)
    print("Wrote " + result["path"])
    print(" Sheets: " + ", ".join(result["sheets"]))
    print(" History: " + str(result["n_hist"]) + " yr | Forecast: " + str(result["n_fcst"]) + " yr")
    chk_str = "ALL PASS" if result["checks_pass"] else "FAILED"
    print(" Checks: " + chk_str + " (" + str(result["n_checks"]) + " items)")

if __name__ == "__main__":
    raise SystemExit(main())
