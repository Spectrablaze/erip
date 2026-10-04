#!/usr/bin/env python3
"""
build_model.py - turn `model_input.json` into a live three-statement model.

    python build_model.py --in data/model_input.json --outdir data/

Writes:
    model.xlsx   a linked workbook. Historical cells and drivers are hardcoded
                 (blue); every projected cell is a real Excel formula, so
                 changing a driver on the Drivers sheet recalculates the whole
                 model in Excel.
    model.json   the same model solved in Python, for the report pipeline and
                 for check_model.py.

Both come out of `rows.py` through `engine.py`, so the workbook and the JSON are
the same model by construction.

Interest is struck on average debt and the revolver is funded from a cash flow
that contains that interest, so the model is deliberately circular. Python
iterates to a fixed point; the workbook is saved with iterative calculation on.
"""
from __future__ import annotations

import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import rows as R          # noqa: E402
from engine import Formula, col_letter, solve  # noqa: E402

DATA_COL0 = 1        # data starts in column B
HDR_ROW = 4          # the period header row
FIRST_ROW = 5

BLUE = "FF0000FF"    # hardcoded input
GREEN = "FF008000"   # link to another sheet
BLACK = "FF000000"   # formula on this sheet
GREY = "FF808080"


# ------------------------------------------------------------------- assembly
def expand(value, n, key, notes):
    """A driver value -> one value per forecast year."""
    if isinstance(value, list):
        v = [float(x) for x in value][:n]
        if len(v) < n:
            notes.append(f"driver '{key}': {len(v)} year(s) supplied for a {n}y "
                         f"forecast, held flat at {v[-1] if v else 0} thereafter")
            v += [v[-1] if v else 0.0] * (n - len(v))
        return v
    return [float(value)] * n


def build_grid(doc, notes):
    hist, fcst = doc["hist_periods"], doc["fcst_periods"]
    nh, nf = len(hist), len(fcst)
    ncols = nh + nf
    actuals = doc.get("actuals", {})
    drivers = dict(doc.get("drivers", {}))

    R.apply_options(doc.get("options", {}).get("dep_basis", "net_block"))
    keys = [r.key for r in R.ROWS if r.kind != "head"]

    missing_actuals = []
    fixed, formulas = {}, {}

    for row in R.ROWS:
        if row.kind == "head":
            continue
        hist_f = Formula(row.hist) if row.hist else (Formula(row.fcst) if row.fcst else None)
        fcst_f = Formula(row.fcst) if row.fcst else None

        for j in range(nh):
            if row.kind == "actual":
                v = (actuals.get(row.key) or [None] * nh)
                v = v[j] if j < len(v) else None
                if v is None:
                    missing_actuals.append((row.key, hist[j]))
                    fixed[(row.key, j)] = 0.0
                else:
                    fixed[(row.key, j)] = float(v)
            elif hist_f is not None:
                formulas[(row.key, j)] = hist_f
            else:
                fixed[(row.key, j)] = 0.0

        for j in range(nh, ncols):
            if row.kind == "input":
                fixed[(row.key, j)] = 0.0        # filled after the first solve
            elif fcst_f is not None:
                formulas[(row.key, j)] = fcst_f
            else:
                fixed[(row.key, j)] = 0.0

    # ---- pass 1: solve the history so the drivers can be back-solved
    grid, iters, resid = solve(keys, ncols, fixed, formulas)

    # ---- fill every driver: supplied value, else the trailing historical average
    avg_years = int(doc.get("options", {}).get("avg_years", 3))
    used, held = {}, []
    for key in R.DRIVER_KEYS:
        if key in drivers:
            vals = expand(drivers[key], nf, key, notes)
            src = doc.get("driver_sources", {}).get(key, "assumptions")
        else:
            lo = max(1, nh - avg_years) if nh > 1 else 0
            sample = [grid[key][j] for j in range(lo, nh)]
            sample = [x for x in sample if x == x]
            mean = sum(sample) / len(sample) if sample else 0.0
            vals = [mean] * nf
            src = f"{len(sample)}y historical average"
            held.append(key)
        used[key] = {"values": vals, "source": src}
        for i, v in enumerate(vals):
            fixed[(key, nh + i)] = v

    if held:
        notes.append("held at the trailing historical average (no value in "
                     "assumptions.json or overrides): " + ", ".join(held))

    # ---- pass 2: solve the whole model
    grid, iters, resid = solve(keys, ncols, fixed, formulas)
    if resid > 1e-4:
        notes.append(f"the circular solve stopped at a residual of {resid:.6g} after "
                     f"{iters} passes - inspect the debt and interest schedule")

    by_row = {}
    for key, period in missing_actuals:
        by_row.setdefault(key, []).append(period)
    absent = [k for k, ps in by_row.items() if len(ps) == nh]
    partial = {k: ps for k, ps in by_row.items() if len(ps) < nh}
    if absent:
        notes.append("not in the source at all, held at zero for every year: "
                     + ", ".join(sorted(absent)))
    for key, ps in sorted(partial.items()):
        notes.append(f"no actual for '{key}' in {', '.join(ps)}: set to zero")

    return {"keys": keys, "grid": grid, "fixed": fixed, "formulas": formulas,
            "ncols": ncols, "nh": nh, "nf": nf, "iters": iters, "resid": resid,
            "drivers_used": used, "missing_actuals": missing_actuals}


# --------------------------------------------------------------------- checks
def run_checks(M, doc):
    grid, nh, ncols = M["grid"], M["nh"], M["ncols"]
    cols = doc["hist_periods"] + doc["fcst_periods"]
    out = {"pass": True, "items": []}

    def add(name, label, ok, worst, period, kind, detail):
        out["items"].append({"check": name, "label": label, "ok": ok,
                             "worst": round(worst, 4), "period": period,
                             "severity": kind, "detail": detail})
        if not ok and kind == "blocking":
            out["pass"] = False

    scale = max((abs(grid["total_assets"][j]) for j in range(ncols)), default=1.0) or 1.0
    tol = max(0.01, scale * 1e-6)

    # The balance check is the one with real teeth in the forecast: it is the only
    # one that is not true by construction, and it fails the moment a change to
    # rows.py stops the cash flow statement fully articulating the balance sheet.
    for key in R.ZERO_CHECKS:
        vals = [(abs(grid[key][j]), cols[j]) for j in range(nh, ncols)]
        worst, per = max(vals) if vals else (0.0, "-")
        structural = key != "chk_balance"
        add(key, R.BY_KEY[key].label + " (forecast)", worst <= tol, worst, per,
            "blocking",
            "an identity by construction - it fails only if the model definition "
            "itself breaks" if structural
            else "articulation of the cash flow statement into the balance sheet; "
                 "this is the check that carries the model")

    # On history the same residuals are computed against reported figures, so they
    # stop being identities and start being statements about the source data.
    for key, why in [
        ("chk_reserves", "reserves moved by more than profit less dividends - "
                         "look for OCI, a buyback, a share issue or a restatement"),
        ("chk_fa_roll", "the fixed asset roll does not close on the reported net "
                        "block once real capex is supplied"),
    ]:
        vals = [(abs(grid[key][j]), cols[j]) for j in range(1, nh)]
        worst, per = max(vals) if vals else (0.0, "-")
        add(key + "_hist", R.BY_KEY[key].label + " (history)",
            worst <= max(tol, scale * 0.002), worst, per, "advisory", why)

    # Reconstructed against reported cash from operations - an accrual-quality read,
    # not a defect.
    gaps = [(abs(grid["cfo_gap"][j]), cols[j]) for j in range(1, nh)
            if grid["rep_cfo"][j]]
    if gaps:
        worst, per = max(gaps)
        ref = max(abs(grid["rep_cfo"][j]) for j in range(nh)) or 1.0
        add("cfo_reconstruction", "Reconstructed CFO close to reported CFO",
            worst <= 0.25 * ref, worst, per, "advisory",
            "a persistent gap means profit is not converting to operating cash the "
            "way the driver model assumes")

    for key in R.NONNEG_CHECKS:
        vals = [(grid[key][j], cols[j]) for j in range(nh, ncols)]
        worst, per = min(vals) if vals else (0.0, "-")
        add(key, R.BY_KEY[key].label, worst >= -tol, worst, per, "blocking",
            "a negative here means the model is funding itself with something "
            "it does not have")

    # Historical balance: the source should already tie. If it does not, the
    # ingest mapping is wrong, not the forecast.
    hv = [(abs(grid["chk_balance"][j]), cols[j]) for j in range(nh)]
    worst, per = max(hv) if hv else (0.0, "-")
    add("hist_balance", "Historical balance sheet ties to the source",
        worst <= tol, worst, per, "blocking",
        "history comes straight from the source; a break here is an ingest problem")

    # Advisory drift checks.
    for key, lo, hi, label in [
        ("tax_rate", 0.0, 0.50, "Effective tax rate"),
        ("ebitda_margin", -0.20, 0.70, "EBITDA margin"),
        ("rev_growth", -0.30, 0.60, "Revenue growth"),
        ("payout_ratio", 0.0, 1.50, "Dividend payout"),
    ]:
        vals = [(grid[key][j], cols[j]) for j in range(nh, ncols)]
        bad = [(v, c) for v, c in vals if v < lo or v > hi]
        add(f"band_{key}", f"{label} inside a plausible band", not bad,
            (bad[0][0] if bad else 0.0), (bad[0][1] if bad else "-"), "advisory",
            f"expected between {lo:.0%} and {hi:.0%}" if key != "payout_ratio"
            else "expected between 0% and 150%")

    # Forecast drivers far from the historical range. The tolerance is the wider
    # of the historical spread and 15% of the historical level, so a driver that
    # has barely moved for a decade does not flag on the third decimal place.
    for key in R.DRIVER_KEYS:
        if (R.BY_KEY[key].hist or "").strip() == "0":
            continue        # no history to compare against by construction
        h = [grid[key][j] for j in range(1, nh)]
        f = [grid[key][j] for j in range(nh, ncols)]
        if not h or not f or (min(h) == 0 and max(h) == 0):
            continue
        lo, hi = min(h), max(h)
        span = max(hi - lo, 0.15 * max(abs(lo), abs(hi)))
        if span == 0:
            continue
        off = [(v, cols[nh + i]) for i, v in enumerate(f)
               if v < lo - span or v > hi + span]
        if off:
            add(f"drift_{key}", f"{R.BY_KEY[key].label} vs its history", False,
                off[0][0], off[0][1], "advisory",
                f"history ran {lo:.4g} to {hi:.4g}; the forecast steps outside that "
                "by more than its full historical range")

    return out


# ---------------------------------------------------------------------- Excel
def write_xlsx(path, doc, M, checks, notes):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.formatting.rule import CellIsRule
    except ImportError:
        sys.exit("openpyxl is required to write the workbook: pip install openpyxl")

    cols = doc["hist_periods"] + doc["fcst_periods"]
    nh, ncols = M["nh"], M["ncols"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    info = wb.create_sheet("Model info")
    sheets = {name: wb.create_sheet(name) for name in R.SHEETS}

    # ---- lay out row positions first so cross-sheet references can be emitted
    pos, cursor = {}, {name: FIRST_ROW for name in R.SHEETS}
    layout = {name: [] for name in R.SHEETS}
    for row in R.ROWS:
        r = cursor[row.sheet]
        layout[row.sheet].append((row, r))
        if row.kind != "head":
            pos[row.key] = (row.sheet, r)
        cursor[row.sheet] = r + 1

    def ref(key, lag, cur_sheet, j):
        if key not in pos:
            return "0"
        sheet, xr = pos[key]
        c = j - lag
        if c < 0:
            return "0"
        cell = f"{col_letter(DATA_COL0 + c)}{xr}"
        return cell if sheet == cur_sheet else f"'{sheet}'!{cell}"

    thin = Side(style="thin", color="FFBFBFBF")
    hdr_fill = PatternFill("solid", fgColor="FF1F3864")
    fcst_fill = PatternFill("solid", fgColor="FFFFF7E6")
    sec_fill = PatternFill("solid", fgColor="FFEDEDED")
    inp_fill = PatternFill("solid", fgColor="FFFFFDE7")

    for name, sheet in sheets.items():
        sheet.sheet_view.showGridLines = False
        sheet["A1"] = f"{doc['company']} - {name}"
        sheet["A1"].font = Font(bold=True, size=14, color="FF1F3864")
        sheet["A2"] = (f"{doc['currency']} unless stated  |  built "
                       f"{doc.get('built', '')}  |  blue = input, black = formula, "
                       "green = link to another sheet")
        sheet["A2"].font = Font(size=9, color=GREY, italic=True)

        sheet.cell(HDR_ROW, 1, "").font = Font(bold=True)
        for j, label in enumerate(cols):
            c = sheet.cell(HDR_ROW, DATA_COL0 + 1 + j, label)
            c.font = Font(bold=True, color="FFFFFFFF")
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center")
        lbl = sheet.cell(HDR_ROW, 1, "Actuals" + " " * 4)
        lbl.font = Font(bold=True, color="FFFFFFFF")
        lbl.fill = hdr_fill

        for row, xr in layout[name]:
            a = sheet.cell(xr, 1, ("    " * row.indent) + row.label)
            if row.kind == "head":
                a.font = Font(bold=True, size=10, color="FF1F3864")
                for j in range(ncols + 1):
                    sheet.cell(xr, 1 + j).fill = sec_fill
                continue
            a.font = Font(bold=row.bold, size=10)
            a.alignment = Alignment(indent=row.indent)

            for j in range(ncols):
                cell = sheet.cell(xr, DATA_COL0 + 1 + j)
                cell.number_format = R.FMT.get(row.fmt, R.FMT["cur"])
                if row.rule == "top":
                    cell.border = Border(top=thin)
                if j >= nh and name != R.CHK:
                    cell.fill = fcst_fill

                if (row.key, j) in M["formulas"]:
                    body = M["formulas"][(row.key, j)].excel(
                        lambda k, lag, _s=name, _j=j: ref(k, lag, _s, _j))
                    cell.value = "=" + body
                    cell.font = Font(bold=row.bold, size=10,
                                     color=GREEN if "!" in body else BLACK)
                else:
                    cell.value = round(M["fixed"].get((row.key, j), 0.0), 6)
                    cell.font = Font(bold=row.bold, size=10, color=BLUE)
                    if row.kind == "input" and j >= nh:
                        cell.fill = inp_fill

        sheet.column_dimensions["A"].width = 46
        for j in range(ncols):
            sheet.column_dimensions[col_letter(DATA_COL0 + 1 + j)].width = 13
        sheet.freeze_panes = sheet.cell(FIRST_ROW, DATA_COL0 + 1)

    # ---- highlight a broken check without needing to read the number
    chk = sheets[R.CHK]
    first, last = FIRST_ROW, cursor[R.CHK] - 1
    rng = (f"{col_letter(DATA_COL0 + 1)}{first}:"
           f"{col_letter(DATA_COL0 + ncols)}{last}")
    red = PatternFill("solid", fgColor="FFFFC7CE")
    chk.conditional_formatting.add(rng, CellIsRule(
        operator="notBetween", formula=["-0.01", "0.01"], fill=red))

    # ---- model info
    info.sheet_view.showGridLines = False
    info["A1"] = f"{doc['company']} - three-statement model"
    info["A1"].font = Font(bold=True, size=16, color="FF1F3864")
    line = 3

    def put(label, value, bold=False):
        nonlocal line
        info.cell(line, 1, label).font = Font(bold=True, size=10)
        info.cell(line, 2, value).font = Font(bold=bold, size=10)
        line += 1

    put("Currency", doc["currency"])
    put("Sector", doc.get("sector") or "not resolved")
    put("History", f"{cols[0]} to {cols[nh - 1]}  ({nh} years)")
    put("Forecast", f"{cols[nh]} to {cols[-1]}  ({M['nf']} years)")
    put("Depreciation basis", doc.get("options", {}).get("dep_basis", "net_block"))
    put("Circular solve", f"{M['iters']} passes, residual {M['resid']:.3g}")
    put("Checks", "ALL PASS" if checks["pass"] else "FAILED - see the Checks sheet",
        bold=True)
    line += 1

    info.cell(line, 1, "Integrity checks").font = Font(bold=True, size=12,
                                                       color="FF1F3864")
    line += 1
    for it in checks["items"]:
        info.cell(line, 1, ("OK   " if it["ok"] else "FAIL ") + it["label"]).font = \
            Font(size=10, color=BLACK if it["ok"] else "FFC00000")
        info.cell(line, 2, it["severity"])
        info.cell(line, 3, it["worst"])
        info.cell(line, 4, it["period"])
        line += 1

    if doc.get("reconciliation"):
        line += 1
        info.cell(line, 1, "Screener vs annual report").font = Font(
            bold=True, size=12, color="FF1F3864")
        line += 1
        for h, t in enumerate(["line", "period", "screener", "annual report", "diff %"]):
            info.cell(line, 1 + h, t).font = Font(bold=True, size=9)
        line += 1
        for d in doc["reconciliation"]:
            for h, t in enumerate([d["row"], d["period"], d["screener"],
                                   d["annual_report"], d["diff_pct"]]):
                info.cell(line, 1 + h, t).font = Font(size=10)
            line += 1

    if doc.get("citations"):
        line += 1
        info.cell(line, 1, "Annual-report citations").font = Font(
            bold=True, size=12, color="FF1F3864")
        line += 1
        for k, c in doc["citations"].items():
            src = c.get("source", "") if isinstance(c, dict) else str(c)
            page = c.get("page", "") if isinstance(c, dict) else ""
            info.cell(line, 1, k).font = Font(size=10)
            info.cell(line, 2, f"{src}  p.{page}").font = Font(size=10)
            line += 1

    if notes:
        line += 1
        info.cell(line, 1, "Notes").font = Font(bold=True, size=12, color="FF1F3864")
        line += 1
        for nline in notes:
            info.cell(line, 1, "- " + nline).font = Font(size=10)
            line += 1

    info.column_dimensions["A"].width = 62
    for c in "BCDE":
        info.column_dimensions[c].width = 18

    # ---- the model is circular by design; let Excel resolve it the same way
    try:
        wb.calculation.iterate = True
        wb.calculation.iterateCount = 200
        wb.calculation.iterateDelta = 1e-6
        wb.calculation.fullCalcOnLoad = True
    except AttributeError:
        notes.append("could not set iterative calculation: switch it on in Excel "
                     "under File > Options > Formulas > Enable iterative calculation")

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


# ----------------------------------------------------------------------- JSON
def write_json(path, doc, M, checks, notes):
    cols = doc["hist_periods"] + doc["fcst_periods"]
    nh = M["nh"]
    out = {
        "company": doc["company"],
        "currency": doc["currency"],
        "sector": doc.get("sector"),
        "built": doc.get("built"),
        "periods": {"all": cols, "hist": doc["hist_periods"],
                    "fcst": doc["fcst_periods"], "n_hist": nh},
        "options": doc.get("options", {}),
        "solve": {"iterations": M["iters"], "residual": M["resid"]},
        "rows": {},
        "drivers": M["drivers_used"],
        "checks": checks,
        "reconciliation": doc.get("reconciliation", []),
        "citations": doc.get("citations", {}),
        "notes": notes,
    }
    for row in R.ROWS:
        if row.kind == "head":
            continue
        out["rows"][row.key] = {
            "label": row.label, "sheet": row.sheet, "kind": row.kind,
            "unit": row.fmt, "note": row.note,
            "values": [round(v, 6) for v in M["grid"][row.key]],
        }

    # A compact forecast bridge for the valuation layer downstream.
    def f(key):
        return [round(M["grid"][key][j], 4) for j in range(nh, M["ncols"])]

    out["bridge"] = {
        "periods": doc["fcst_periods"],
        "revenue": f("revenue"), "ebitda": f("ebitda"), "ebit": f("ebit"),
        "depreciation": f("dep_charge"), "capex": f("capex"),
        "change_in_nwc": f("d_nwc"), "tax_rate": f("tax_rate"),
        "fcff": f("s_fcff"), "fcfe": f("s_fcfe"),
        "net_debt_close": f("s_net_debt"), "shares": f("shares"),
        "eps": f("eps"), "bvps": f("bvps"),
    }
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=1)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--in", dest="inp", default="data/model_input.json")
    ap.add_argument("--outdir", default="data")
    ap.add_argument("--xlsx", default=None, help="override the workbook path")
    ap.add_argument("--strict", action="store_true",
                    help="exit non-zero when a blocking check fails")
    ap.add_argument("--no-xlsx", action="store_true")
    args = ap.parse_args()

    with open(args.inp, encoding="utf-8") as f:
        doc = json.load(f)

    notes = list(doc.get("notes", []))
    M = build_grid(doc, notes)
    checks = run_checks(M, doc)

    jpath = os.path.join(args.outdir, "model.json")
    write_json(jpath, doc, M, checks, notes)
    xpath = args.xlsx or os.path.join(args.outdir, "model.xlsx")
    if not args.no_xlsx:
        write_xlsx(xpath, doc, M, checks, notes)

    cols = doc["hist_periods"] + doc["fcst_periods"]
    g = M["grid"]
    print(f"{doc['company']}  ({doc['currency']})")
    print(f"  history  {cols[0]} .. {cols[M['nh'] - 1]}"
          f"   forecast {cols[M['nh']]} .. {cols[-1]}")
    print(f"  circular solve converged in {M['iters']} passes "
          f"(residual {M['resid']:.3g})\n")

    head = "  " + "line".ljust(22) + "".join(c.rjust(12) for c in cols[M["nh"]:])
    print(head)
    for key in ("revenue", "ebitda", "ebit", "pat", "eps", "s_fcff",
                "cash", "debt_total", "s_net_debt"):
        vals = "".join(f"{g[key][j]:>12,.1f}" for j in range(M["nh"], M["ncols"]))
        print("  " + R.BY_KEY[key].label[:22].ljust(22) + vals)

    print()
    bad = [i for i in checks["items"] if not i["ok"]]
    if not bad:
        print(f"  checks   all {len(checks['items'])} pass")
    else:
        for i in bad:
            print(f"  {i['severity'].upper():<9} {i['label']}: "
                  f"{i['worst']:,.4g} in {i['period']}  - {i['detail']}")
    for n in notes:
        print(f"  ! {n}")

    print(f"\nwrote {jpath}")
    if not args.no_xlsx:
        print(f"wrote {xpath}")
    return 1 if (args.strict and not checks["pass"]) else 0


if __name__ == "__main__":
    raise SystemExit(main())
